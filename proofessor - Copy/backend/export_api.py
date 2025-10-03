
from fastapi import APIRouter, Body
from pydantic import BaseModel
from fastapi.responses import FileResponse
import os
import uuid
from docx import Document
from fpdf import FPDF

from dotenv import load_dotenv

load_dotenv()

EXPORT_DIR = os.getenv("EXPORT_DIR", "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)

router = APIRouter()

class ExportRequest(BaseModel):
    text: str
    format: str  # 'docx', 'pdf', 'latex'

@router.post("/export_draft/")
async def export_draft(req: ExportRequest):
    filename = f"draft_{uuid.uuid4().hex[:8]}.{req.format}"
    path = os.path.join(EXPORT_DIR, filename)
    if req.format == "docx":
        doc = Document()
        doc.add_paragraph(req.text)
        doc.save(path)
    elif req.format == "pdf":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        for line in req.text.split("\n"):
            pdf.cell(200, 10, txt=line, ln=1)
        pdf.output(path)
    elif req.format == "latex":
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"""\\documentclass{{article}}\n\\begin{{document}}\n{req.text}\n\\end{{document}}""")
    else:
        return {"error": "unsupported format"}
    return {"filename": filename}

@router.get("/download_export/{filename}")
async def download_export(filename: str):
    path = os.path.join(EXPORT_DIR, filename)
    if not os.path.exists(path):
        return {"error": "not found"}
    return FileResponse(path, filename=filename)
